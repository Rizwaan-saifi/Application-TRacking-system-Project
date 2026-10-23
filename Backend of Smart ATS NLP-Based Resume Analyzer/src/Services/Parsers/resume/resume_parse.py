import pdfplumber
from docx import Document

# ---------------------------------------------------------
# STEP 1: RAW TEXT EXTRACTORS (PDF & DOCX)
# ---------------------------------------------------------

def extract_raw_text_from_pdf(pdf_path: str) -> str:
    """PDF file se line-by-line raw text extract karta hai."""
    lines = []
    try:
        with pdfplumber.open(pdf_path) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    lines.extend(text.split("\n"))
        return "\n".join(lines)
    except Exception as e:
        raise Exception(f"Error reading PDF file: {str(e)}")

def extract_raw_text_from_docx(docx_path: str) -> str:
    """DOCX file se paragraphs aur tables ka raw text extract karta hai."""
    lines = []
    try:
        doc = Document(docx_path)
        # Paragraphs extraction
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                lines.append(paragraph.text.strip())
        
        # Tables extraction (multi-column layouts ke liye)
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    lines.append(" | ".join(row_text))
                    
        return "\n".join(lines)
    except Exception as e:
        raise Exception(f"Error reading DOCX file: {str(e)}")